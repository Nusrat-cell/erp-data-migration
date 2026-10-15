import shutil
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from migrate import pipeline, report, sample

ROOT = Path(__file__).resolve().parent.parent


class Parsing(unittest.TestCase):
    def test_amounts_in_both_formats(self):
        self.assertEqual(pipeline.parse_amount("1.234,50"), 123450)
        self.assertEqual(pipeline.parse_amount("1234.50"), 123450)
        self.assertEqual(pipeline.parse_amount("-11.191,87"), -1119187)
        self.assertEqual(pipeline.parse_amount("0,05"), 5)
        self.assertIsNone(pipeline.parse_amount(""))
        self.assertIsNone(pipeline.parse_amount("abc"))

    def test_no_float_rounding(self):
        # 0.1 + 0.2 style errors must not appear when summing many small amounts.
        self.assertEqual(sum(pipeline.parse_amount("0,10") for _ in range(3)), 30)

    def test_dates(self):
        self.assertEqual(pipeline.parse_date("14-03-2025"), "2025-03-14")
        self.assertEqual(pipeline.parse_date("2025-03-14"), "2025-03-14")
        self.assertIsNone(pipeline.parse_date("31-02-2025"))
        self.assertIsNone(pipeline.parse_date("March 14"))

    def test_countries_and_names(self):
        self.assertEqual(pipeline.country_code(" Danmark "), "DK")
        self.assertEqual(pipeline.country_code("Sverige"), "SE")
        self.assertIsNone(pipeline.country_code("Atlantis"))
        self.assertEqual(pipeline.match_name("Lyngby Trading A/S"), pipeline.match_name("Lyngby Trading AS"))


class TrialRun(unittest.TestCase):
    """First trial migration on the sample export: every planted problem is caught."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        sample.generate(cls.tmp / "legacy")
        cls.r = pipeline.run(cls.tmp / "legacy")
        cls.rules = cls.r.issues.groupby("record").rule.apply(set).to_dict()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp)

    def test_status_is_no_go_with_balanced_numbers(self):
        self.assertFalse(self.r.go)
        self.assertTrue(self.r.balanced)
        self.assertEqual(self.r.blocking, 5)

    def test_exact_duplicate_removed_once(self):
        self.assertIn("CUST-DUP-EXACT", self.rules["C1007"])
        self.assertEqual((self.r.customers["No."] == "C1007").sum(), 1)

    def test_same_company_merged_and_entry_moved(self):
        self.assertIn("CUST-DUP-MATCH", self.rules["C1031"])
        self.assertNotIn("C1031", set(self.r.customers["No."]))
        self.assertEqual(self.r.dedupe_map.to_dict("records"), [{"legacy_no": "C1031", "surviving_no": "C1012"}])
        moved = self.r.invoices[self.r.invoices["Document No."] == "INV-5012"]
        self.assertEqual(moved["Customer No."].item(), "C1012")

    def test_missing_vat_holds_customer_and_its_entries(self):
        self.assertIn("CUST-VAT", self.rules["C1019"])
        self.assertNotIn("C1019", set(self.r.customers["No."]))
        self.assertIn("INV-CUST-HELD", self.rules["INV-5019"])

    def test_bad_email_blanked_not_rejected(self):
        row = self.r.customers[self.r.customers["No."] == "C1023"]
        self.assertEqual(row["E-Mail"].item(), "")

    def test_countries_and_dates_normalised(self):
        self.assertTrue(set(self.r.customers["Country/Region Code"]) <= {"DK", "SE", "NO", "DE"})
        self.assertTrue(self.r.customers["Created Date"].str.match(r"^\d{4}-\d{2}-\d{2}$").all())

    def test_duplicate_sku_held(self):
        self.assertIn("ITEM-DUP-SKU", self.rules["ITM-207"])
        self.assertNotIn("ITM-207", set(self.r.items["No."]))

    def test_invoice_problems(self):
        self.assertIn("INV-ORPHAN", self.rules["INV-5040"])
        self.assertIn("INV-SIGN", self.rules["INV-5033"])
        fixed = self.r.invoices[self.r.invoices["Document No."] == "INV-5027"]
        self.assertEqual(fixed["Currency Code"].item(), "DKK")

    def test_every_planted_problem_has_a_rule(self):
        expected = {"CUST-DUP-EXACT", "CUST-DUP-MATCH", "CUST-VAT", "CUST-EMAIL", "CUST-COUNTRY",
                    "ITEM-DUP-SKU", "INV-ORPHAN", "INV-SIGN", "INV-CURRENCY"}
        self.assertTrue(expected <= set(self.r.issues.rule))
        self.assertEqual(len(sample.PLANTED), 11)

    def test_amounts_reconcile_to_the_cent(self):
        t = self.r.totals
        self.assertTrue(t.balanced.all())
        dkk = t[t.currency == "DKK"].iloc[0]
        self.assertAlmostEqual(dkk.source, dkk.loaded + dkk.held, places=2)


class AfterDecisions(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        sample.generate(self.tmp / "legacy")

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_decisions_turn_no_go_into_go(self):
        decisions = pipeline.load_decisions(ROOT / "decisions.example.yaml")
        r = pipeline.run(self.tmp / "legacy", decisions=decisions)
        self.assertTrue(r.go, r.issues[r.issues.severity == "blocking"])
        self.assertEqual(int(r.reconciliation.dropped_by_decision.sum()), 2)
        self.assertEqual(r.totals.set_index("currency").loc["DKK", "dropped_by_decision"], 12500.0)
        self.assertEqual((r.issues.rule == "DECISION").sum(), 4)

    def test_decision_that_matches_nothing_fails_loudly(self):
        bad = [{"entity": "items", "record": "ITM-999", "action": "drop"}]
        with self.assertRaises(ValueError):
            pipeline.run(self.tmp / "legacy", decisions=bad)

    def test_run_twice_gives_same_output(self):
        a = pipeline.run(self.tmp / "legacy")
        b = pipeline.run(self.tmp / "legacy")
        pd.testing.assert_frame_equal(a.invoices, b.invoices)
        pd.testing.assert_frame_equal(a.issues, b.issues)


class Report(unittest.TestCase):
    def test_workbook_and_load_files(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            sample.generate(tmp / "legacy")
            r = pipeline.run(tmp / "legacy")
            path = report.write(r, tmp / "out")
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames[:5], ["Summary", "Issues", "Reconciliation", "OpenAmountTotals", "DedupeMap"])
            self.assertTrue(wb["Summary"]["B2"].value.startswith("NO-GO"))
            load = pd.read_csv(tmp / "out" / "load_Customers.csv", sep=";", dtype=str)
            self.assertEqual(list(load.columns)[:2], ["No.", "Name"])
            self.assertEqual(len(load), len(r.customers))
        finally:
            shutil.rmtree(tmp)


if __name__ == "__main__":
    unittest.main()
