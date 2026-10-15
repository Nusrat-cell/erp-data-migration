"""Write the migration result as load files and one Excel workbook for the business."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER = PatternFill("solid", fgColor="0B6E5C")
RED = PatternFill("solid", fgColor="F8D7DA")
AMBER = PatternFill("solid", fgColor="FFF3CD")
GREEN = PatternFill("solid", fgColor="D4EDDA")


def _sheet(wb, title, df, severity_col=None):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([v.item() if hasattr(v, "item") else v for v in row])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    if severity_col and severity_col in df.columns:
        idx = list(df.columns).index(severity_col) + 1
        for r in range(2, ws.max_row + 1):
            sev = ws.cell(r, idx).value
            fill = RED if sev == "blocking" else AMBER if sev == "warning" else None
            if fill:
                for c in ws[r]:
                    c.fill = fill
    return ws


def summary_lines(result):
    issues = result.issues
    counts = issues.severity.value_counts().to_dict() if len(issues) else {}
    status = "GO" if result.go else "NO-GO"
    reason = ("all records reconcile and nothing is blocking" if result.go else
              f"{result.blocking} blocking issue(s) need a decision" if result.balanced else
              "reconciliation does not balance")
    return [
        ("Go-live status", f"{status}: {reason}"),
        ("Customers loaded", int(result.customers.shape[0])),
        ("Items loaded", int(result.items.shape[0])),
        ("Open entries loaded", int(result.invoices.shape[0])),
        ("Blocking issues", counts.get("blocking", 0)),
        ("Warnings (fixed or flagged)", counts.get("warning", 0)),
        ("Info (automatic fixes)", counts.get("info", 0)),
        ("Record counts balance", "yes" if result.reconciliation.balanced.all() else "NO"),
        ("Open amounts balance per currency", "yes" if result.totals.balanced.all() else "NO"),
    ]


def write(result, out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    m = result.mapping
    loads = {m["customers"]["target"]: result.customers, m["items"]["target"]: result.items,
             m["open_invoices"]["target"]: result.invoices}
    for name, df in loads.items():
        df.to_csv(out / f"load_{name}.csv", index=False, sep=";", encoding="utf-8")
    result.issues.to_csv(out / "issue_log.csv", index=False, sep=";", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Data migration report", ""])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in summary_lines(result):
        ws.append([k, v])
    ws["B2"].fill = GREEN if result.go else RED
    ws["B2"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    ws.append([])
    ws.append(["How to read this", "Blocking rows are held back and listed in Issues with the owner to ask. "
               "Every source record is either loaded, merged, removed as an identical copy or held, "
               "and the Reconciliation sheet proves the numbers add up."])
    ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)

    _sheet(wb, "Issues", result.issues, severity_col="severity")
    _sheet(wb, "Reconciliation", result.reconciliation)
    _sheet(wb, "OpenAmountTotals", result.totals)
    _sheet(wb, "DedupeMap", result.dedupe_map)
    for name, df in loads.items():
        _sheet(wb, name, df)
    path = out / "migration_report.xlsx"
    wb.save(path)
    return path
